import io
import os
from typing import Any

import docx
import pymupdf as fitz

def extract_text_from_pdf(file_content: bytes) -> dict[str, Any]:
    if not file_content:
        raise ValueError("PDF file content is empty.")

    try:
        pdf_stream = io.BytesIO(file_content)
        doc = fitz.open(stream=pdf_stream.read(), filetype="pdf")
    except Exception as exc:
        raise ValueError(f"Failed to open PDF document: {exc}") from exc

    try:
        num_pages = len(doc)
        page_texts: list[str] = []

        for page_idx in range(num_pages):
            page = doc[page_idx]
            extracted = page.get_text()
            if extracted:
                page_texts.append(extracted.strip())

        full_text = "\n\n".join(page_texts).strip()

        avg_chars = len(full_text) / max(num_pages, 1) if num_pages > 0 else 0
        is_scanned = (avg_chars < 30) or (num_pages == 0)

        raw_metadata = doc.metadata or {}
        clean_metadata: dict[str, str] = {
            str(k): str(v) for k, v in raw_metadata.items() if v is not None
        }

        return {
            "text": full_text,
            "num_pages": num_pages,
            "is_scanned": is_scanned,
            "metadata": clean_metadata,
        }
    finally:
        doc.close()

def extract_text_from_docx(file_content: bytes) -> dict[str, Any]:
    if not file_content:
        raise ValueError("DOCX file content is empty.")

    try:
        docx_stream = io.BytesIO(file_content)
        doc = docx.Document(docx_stream)
    except Exception as exc:
        raise ValueError(f"Failed to open DOCX document: {exc}") from exc

    paragraph_texts: list[str] = [
        p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()
    ]
    num_paragraphs = len(doc.paragraphs)

    table_texts: list[str] = []
    for table in doc.tables:
        for row in table.rows:
            row_items: list[str] = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text and (not row_items or cell_text != row_items[-1]):
                    row_items.append(cell_text)
            if row_items:
                table_texts.append(" | ".join(row_items))

    combined_chunks = paragraph_texts + table_texts
    full_text = "\n".join(combined_chunks).strip()

    metadata: dict[str, str] = {}
    if hasattr(doc, "core_properties") and doc.core_properties:
        props = doc.core_properties
        tracked_attrs = [
            "author",
            "category",
            "comments",
            "content_status",
            "created",
            "identifier",
            "keywords",
            "language",
            "last_modified_by",
            "modified",
            "subject",
            "title",
            "version",
        ]
        for attr in tracked_attrs:
            try:
                val = getattr(props, attr, None)
                if val is not None:
                    metadata[attr] = str(val)
            except Exception:
                pass

    return {
        "text": full_text,
        "num_paragraphs": num_paragraphs,
        "metadata": metadata,
    }

def extract_text(file_content: bytes, filename: str) -> dict[str, Any]:
    if not filename or "." not in filename:
        raise ValueError(
            f"Invalid filename '{filename}'. Filename must include an extension."
        )

    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        result = extract_text_from_pdf(file_content)
        result["source_format"] = "pdf"
        return result
    elif ext == ".docx":
        result = extract_text_from_docx(file_content)
        result["source_format"] = "docx"
        return result
    else:
        raise ValueError(
            f"Unsupported file format '{ext}'. Only .pdf and .docx are supported."
        )

def validate_file(
    file_content: bytes, filename: str, max_size_mb: int = 10
) -> tuple[bool, str]:
    if not file_content or len(file_content) == 0:
        return False, "File is empty."

    if not filename or "." not in filename:
        return False, "Filename must include a valid extension."

    ext = os.path.splitext(filename)[1].lower()
    allowed_extensions = {".pdf", ".docx"}
    if ext not in allowed_extensions:
        return (
            False,
            f"Unsupported file extension '{ext}'. Supported formats are: .pdf, .docx",
        )

    max_size_bytes = max_size_mb * 1024 * 1024
    if len(file_content) > max_size_bytes:
        size_mb = len(file_content) / (1024 * 1024)
        return (
            False,
            f"File size ({size_mb:.2f} MB) exceeds maximum allowed size ({max_size_mb} MB).",
        )

    return True, ""
